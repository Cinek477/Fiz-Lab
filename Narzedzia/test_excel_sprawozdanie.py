# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl>=3.1,<4", "numpy>=1.26,<3", "matplotlib>=3.8,<4"]
# ///
"""Testy regresyjne obliczeń i pustych kart: uv run Narzedzia/test_excel_sprawozdanie.py."""
from pathlib import Path
import math
import tempfile
import unittest
import numpy as np
import openpyxl
from excel_schema import SPECS,WORKBOOK_PATHS
from excel_sprawozdanie import Report,load,generate,num,fit,pendulum,young,field,slits,spectrum

ROOT=Path(__file__).resolve().parent.parent


def row(**data):return {'_row':7,**data}


class ExcelTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.out=Path(self.temp.name)
    def tearDown(self):self.temp.cleanup()
    def report(self,code,p):return Report(SPECS[code],p,self.out)

    def test_all_blank_templates(self):
        for code,spec in SPECS.items():
            path=ROOT/WORKBOOK_PATHS[code]
            s,p,t=load(path)
            self.assertEqual(s['code'],code);self.assertFalse(any(t.values()))
            with self.assertRaisesRegex(ValueError,'Brak pomiarów'):generate(path,self.out/code,tex_only=True)
            book=openpyxl.load_workbook(path,read_only=False)
            try:
                self.assertEqual(book.sheetnames,['Ustawienia']+[x['name'] for x in spec['tables']])
                for table in spec['tables']:
                    sheet=book[table['name']]
                    self.assertEqual(sheet.freeze_panes,'A7')
                    self.assertEqual(sheet.max_row,table['rows']+6)
                if code=='71':self.assertIn('A1206',book['Pomiary']['B1206'].value)
            finally:book.close()

    def test_zero_missing_and_invalid(self):
        self.assertEqual(num(0,'zero'),0)
        self.assertEqual(num('0,1','przecinek'),.1)
        for v in [None,'',float('nan'),float('inf'),True,'1 mm']:
            with self.assertRaises(ValueError):num(v,'test')
        with self.assertRaises(ValueError):num(0,'dodatnia',positive=True)
        with self.assertRaises(ValueError):fit([1,1,1],[2,3,4])

    def test_pendulum_known_g(self):
        r=self.report('0',{})
        rows=[row(L_cm=L*100,u_L_cm=.1,T_s=math.sqrt(4*math.pi**2*L/9.81),u_T_s=.001) for L in [.1,.2,.3,.4]]
        pendulum(r,{'Czasy':[],'Grupy':rows})
        self.assertAlmostEqual(r.results['regresja']['g'],9.81,places=10)

    def test_young_lever_zero_and_si(self):
        r=self.report('11',{'g':9.81});rows=[];E=200e9;d=.0005
        for m in [1,2,3,4]:
            dl=m*9.81/(math.pi*d*d/4*E)*1000
            rows.append(row(m_kg=m,z_up_mm=.2+2*dl,z_down_mm=.4+2*dl,fit='tak'))
        sample=row(nr='A',material='stal',l_m=1,u_l_m=.001,d1_mm=.5,d2_mm=.5,d3_mm=.5,u_d_mm=.001,z0_up_mm=.2,z0_down_mm=.4)
        young(r,{'Probki':[sample],'Drut_A':rows,'Drut_B':[]})
        self.assertAlmostEqual(r.results['A']['E_GPa'],200,places=9)

    def test_negative_field_and_no_missing_uncertainty_as_zero(self):
        r=self.report('31A',dict(U=10,V0=0,d_mm=70))
        field(r,{'Profile':[row(s_mm=x,Va=10*x/70,Vb=10*x/70,Vc=10*x/70) for x in [10,20,30]],'Mapa':[]})
        self.assertTrue(np.allclose(r.results['profile']['E_V_m'],-10000/70))
        self.assertNotIn('u_E_V_m',r.results['profile'])

    def test_slit_spacing_and_true_zero_minimum(self):
        period=532e-6*734/.3
        r=self.report('71',dict(lambda_nm=532,L_mm=734,Imax=10,Imin=0))
        slits(r,{'Pomiary':[],'Maksima':[row(m=i,left_mm=6-i*period,right_mm=6+i*period) for i in [1,2,3]]})
        self.assertAlmostEqual(r.results['szczeliny']['d_mm'],.3,places=12)
        self.assertEqual(r.results['kontrast']['widzialnosc'],1)
        self.assertTrue(all(v[5] is None for v in r.results['szczeliny']['pary']))

    def test_spectrum_minutes_wrap_and_correlated_zero(self):
        p=dict(N=600,u_N=0,u_angle_min=1,zero1_deg=350,zero1_min=0)
        # 330° i 10° są po przeciwnych stronach zera 350°: oba odchylenia 20°.
        r=self.report('83',p)
        spectrum(r,{'Odczyty':[row(order=1,id=1,repeat=1,ldeg=330,lmin=0,rdeg=10,rmin=0,barwa='kontrolna')],'Odniesienia':[]})
        line=r.results['widmo'][0]
        self.assertAlmostEqual(line['alpha'],20)
        expected=(1e6/600)*math.sin(math.radians(20))
        self.assertAlmostEqual(line['lambda_nm'],expected)
        expected_u=(1e6/600)*math.cos(math.radians(20))*math.radians(1/60)/math.sqrt(2)
        self.assertAlmostEqual(line['u_lambda_nm'],expected_u)
        with self.assertRaises(ValueError):spectrum(self.report('83',p),{'Odczyty':[row(order=1,id=1,repeat=1,ldeg=330,lmin=60,rdeg=10,rmin=0)],'Odniesienia':[]})

    def test_spectrum_right_deviation_matches_raw_scale(self):
        p=dict(N=570,zero1_deg=0,zero1_min=0)
        common=dict(order=1,id=1,repeat=1,ldeg=14,lmin=42,barwa='niebieska')
        raw=self.report('83',p)
        spectrum(raw,{'Odczyty':[row(**common,rdeg=345,rmin=9)],'Odniesienia':[]})
        converted=self.report('83',p)
        spectrum(converted,{'Odczyty':[row(**common,rdeg=14,rmin=51,uwagi='P: odchylenie od 0;')],'Odniesienia':[]})
        self.assertAlmostEqual(raw.results['widmo'][0]['alpha'],converted.results['widmo'][0]['alpha'])
        self.assertAlmostEqual(raw.results['widmo'][0]['lambda_nm'],converted.results['widmo'][0]['lambda_nm'])


if __name__=='__main__':unittest.main()
