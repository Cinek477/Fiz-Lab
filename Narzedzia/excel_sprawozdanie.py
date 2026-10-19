# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl>=3.1,<4", "numpy>=1.26,<3", "matplotlib>=3.8,<4"]
# ///
"""Excel -> obliczenia, wykresy, LaTeX, paczka Overleaf i opcjonalnie PDF."""
import argparse
import csv
import datetime
import json
import math
import os
from pathlib import Path
import shutil
import subprocess
import sys
import zipfile
import numpy as np
import openpyxl
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.tri as tri
from excel_schema import SPECS


def num(v, where, positive=False, nonnegative=False):
    if v is None or v=='':
        raise ValueError(f'Brak danych: {where}.')
    if isinstance(v,bool):raise ValueError(f'{where}: wpisz liczbę, nie wartość logiczną.')
    try:x=float(v.replace(',','.') if isinstance(v,str) else v)
    except (ValueError,TypeError):raise ValueError(f'{where}: oczekiwano liczby; otrzymano {v!r}.')
    if not math.isfinite(x) or (positive and x<=0) or (nonnegative and x<0):
        raise ValueError(f'{where}: niedopuszczalna wartość {v!r}.')
    return x


def missing(v):return v is None or v==''


def load(path):
    book=openpyxl.load_workbook(path,data_only=False,read_only=True)
    try:
        setup=book['Ustawienia'];code=str(setup['B4'].value)
        if code not in SPECS:raise ValueError('Nieznany kod ćwiczenia w Ustawienia!B4.')
        spec=SPECS[code];params={}
        for i,(key,label,_) in enumerate(spec['fields'],7):
            if setup.cell(i,1).value!=label:raise ValueError(f'Ustawienia!A{i}: zmieniono układ lub nazwę parametru.')
            value=setup.cell(i,2).value
            if isinstance(value,str) and value.startswith('='):raise ValueError(f'Ustawienia!B{i}: wpisz wartość, nie formułę.')
            if isinstance(value,(datetime.date,datetime.datetime)):value=value.strftime('%d.%m.%Y')
            params[key]=value
        tables={}
        for t in spec['tables']:
            sh=book[t['name']]
            if [sh.cell(6,c+1).value for c in range(len(t['headers']))]!=t['headers']:
                raise ValueError(f"{t['name']}: zmieniono nagłówki w wierszu 6.")
            result=[]
            for r,values in enumerate(sh.iter_rows(min_row=7,max_col=len(t['keys']),values_only=True),7):
                row=dict(zip(t['keys'],values));row['_row']=r
                if code=='71' and t['name']=='Pomiary':
                    if missing(row['U_mV']):
                        if not missing(row.get('uwagi')):continue
                        continue
                    if isinstance(row['p_mm'],str) and row['p_mm'].startswith('='):
                        a=next(i for i,f in enumerate(spec['fields'],7) if f[0]=='p_start_mm')
                        b=next(i for i,f in enumerate(spec['fields'],7) if f[0]=='krok_mm')
                        expected=f'=IF(COUNT(\'Ustawienia\'!$B${a},\'Ustawienia\'!$B${b})<2,"",\'Ustawienia\'!$B${a}+(A{r}-1)*\'Ustawienia\'!$B${b})'
                        if row['p_mm']!=expected:raise ValueError(f'Pomiary!B{r}: nieznana formuła. Wpisz położenie jako liczbę.')
                        index=num(row['nr'],f'Pomiary!A{r}',positive=True)
                        row['p_mm']=num(params['p_start_mm'],'początek osi')+(index-1)*num(params['krok_mm'],'krok',positive=True)
                elif code=='83' and t['name']=='Odczyty':
                    if all(missing(row[k]) for k in ['ldeg','lmin','rdeg','rmin']):continue
                elif all(missing(row[k]) for k in t['keys'] if k not in ['nr','id','uwagi','fit']):continue
                for k,v in row.items():
                    if isinstance(v,str) and v.startswith('='):raise ValueError(f"{t['name']}!wiersz {r}: formuły są dozwolone tylko w automatycznej kolumnie położenia 71.")
                result.append(row)
            tables[t['name']]=result
    finally:book.close()
    return spec,params,tables


def esc(value):
    s='' if value is None else str(value)
    chars={'\\':r'\textbackslash{}','&':r'\&','%':r'\%','$':r'\$','#':r'\#','_':r'\_','{':r'\{','}':r'\}','~':r'\textasciitilde{}','^':r'\textasciicircum{}'}
    return ''.join(chars.get(c,c) for c in s).replace('\n',r'\newline ')


def fmt(v):
    return esc(f'{v:.6g}'.replace('.',',')) if isinstance(v,(float,np.floating)) else esc(v)


def fit(x,y,through_zero=False):
    x=np.asarray(x,float);y=np.asarray(y,float)
    if len(x)<(2 if through_zero else 3) or len(set(x))<2:raise ValueError('Regresja wymaga co najmniej 3 punktów (2 przy wyrazie wolnym równym 0) i różnych odciętych.')
    X=x[:,None] if through_zero else np.column_stack([x,np.ones(len(x))])
    beta=np.linalg.lstsq(X,y,rcond=None)[0]
    res=y-X@beta;variance=float(res@res/(len(x)-len(beta)))
    cov=variance*np.linalg.inv(X.T@X)
    return float(beta[0]),0. if through_zero else float(beta[1]),float(np.sqrt(cov[0,0])),res


class Report:
    def __init__(self,spec,params,out):
        self.spec=spec;self.p=params;self.out=out;self.parts=[];self.results={};self.warnings=[]
    def text(self,s):self.parts.append(esc(s)+'\n')
    def raw(self,s):self.parts.append(s+'\n')
    def heading(self,s):self.raw(r'\subsection{'+esc(s)+'}')
    def warn(self,s):self.warnings.append(s)
    def table(self,headers,rows):
        if not rows:return
        n=len(headers)
        width=r'\dimexpr(\linewidth-'+str(2*n)+r'\tabcolsep-'+str(n+1)+r'\arrayrulewidth)/'+str(n)+r'\relax'
        cols='|'+'|'.join(r'>{\raggedright\arraybackslash}p{'+width+'}' for _ in headers)+'|'
        head=r'\rowcolor{pale}'+ ' & '.join(r'\textbf{'+esc(h)+'}' for h in headers)+r'\\\hline'
        self.raw(r'\begingroup\small\begin{longtable}{'+cols+'}'+r'\hline'+head+r'\endfirsthead\hline'+head+r'\endhead')
        for row in rows:self.raw(' & '.join(fmt(v) for v in row)+r'\\\hline')
        self.raw(r'\end{longtable}\endgroup')
    def figure(self,name,caption):
        self.raw(r'\begin{figure}[H]\centering\includegraphics[width=.94\linewidth,height=90mm,keepaspectratio]{'+name+r'.pdf}\caption{'+esc(caption)+r'}\end{figure}')
    def savefig(self,name,caption):
        plt.savefig(self.out/(name+'.pdf'),bbox_inches='tight');plt.close();self.figure(name,caption)
    def write(self):
        meta=[('Autorzy','autorzy'),('Wydział','wydzial'),('Rok','rok'),('Grupa','grupa'),('Zespół','zespol'),('Data wykonania','data_wykonania'),('Data oddania','data_oddania')]
        pre=r'''\documentclass[11pt,a4paper]{article}
\usepackage{fontspec}\setmainfont{texgyretermes-regular.otf}[BoldFont=texgyretermes-bold.otf,ItalicFont=texgyretermes-italic.otf]
\usepackage[polish]{babel}\usepackage{geometry,amsmath,amssymb,graphicx,longtable,array,caption,float}
\usepackage[table]{xcolor}\usepackage[hidelinks]{hyperref}\usepackage{fancyhdr}
\geometry{margin=20mm}\definecolor{pale}{HTML}{EDF2F8}
\setlength{\parindent}{0pt}\setlength{\parskip}{6pt}\renewcommand{\arraystretch}{1.25}\raggedbottom
\pagestyle{fancy}\fancyhf{}\fancyfoot[C]{\thepage}\setlength{\headheight}{15pt}
\fancyhead[L]{Pracownia fizyczna}\fancyhead[R]{Fiz-Lab 1.1}
\begin{document}
'''
        intro=Report(self.spec,self.p,self.out)
        intro.raw(r'{\LARGE\bfseries Ćwiczenie '+esc(self.spec['code'])+' — '+esc(self.spec['title'])+r'\par}\medskip')
        intro.table(['Metryka','Dane'],[(label,self.p.get(key) or '—') for label,key in meta])
        intro.raw(r'\section{Cel ćwiczenia}');intro.text(self.spec['goal'])
        intro.raw(r'\section{Podstawy teoretyczne}')
        intro.raw(r'\[\boxed{\begin{gathered}'+self.spec['theory']+r'\end{gathered}}\]')
        intro.raw(r'\section{Stanowisko i przebieg pomiarów}')
        intro.text(self.p.get('przebieg') or '[Uzupełnij opis przebiegu w arkuszu Ustawienia.]')
        if self.p.get('przyrzady'):intro.text(self.p['przyrzady'])
        fields={k:label for k,label,_ in self.spec['fields'][10:]}
        intro.table(['Parametr','Wpis'],[(fields[k],v) for k,v in self.p.items() if k in fields and not missing(v)])
        intro.raw(r'\section{Wyniki i obliczenia}')
        ending=Report(self.spec,self.p,self.out);ending.raw(r'\section{Niepewności i ograniczenia}')
        if self.warnings:
            for w in dict.fromkeys(self.warnings):ending.text(w)
        ending.text('Wpisy u(...) w arkuszu są niepewnościami standardowymi. Nie są automatycznie utożsamiane z działką ani błędem granicznym przyrządu. Niepewność rozszerzoną podaj dopiero po przyjęciu współczynnika k.')
        ending.raw(r'\section{Wnioski}');ending.table(['Wnioski i interpretacja'],[[self.p.get('wnioski') or '[Uzupełnij wnioski w arkuszu Ustawienia na podstawie otrzymanych wyników.]']])
        tex=pre+'\n'.join(intro.parts+self.parts+ending.parts)+r'\end{document}'
        (self.out/'sprawozdanie.tex').write_text(tex,encoding='utf-8')
        (self.out/'obliczenia.json').write_text(json.dumps(self.results,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')


def require(p,key,positive=False):return num(p.get(key),f'Ustawienia: {key}',positive=positive)
def uncertainty(p,key):
    return None if missing(p.get(key)) else num(p[key],f'u: {key}',nonnegative=True)
def values(rows,key):return np.array([num(r.get(key),f'wiersz {r["_row"]}: {key}') for r in rows])
def axes(xlabel,ylabel):
    fig,ax=plt.subplots(figsize=(8,4.2),layout='constrained');ax.set(xlabel=xlabel,ylabel=ylabel);ax.grid(alpha=.2);return fig,ax


def pendulum(r,t):
    p=r.p
    if t['Czasy']:
        times=values(t['Czasy'],'t_s');n=require(p,'n',True);L=require(p,'L_m',True)
        if n!=int(n) or len(times)<2 or np.any(times<=0):raise ValueError('Wahadło: n musi być całkowite, potrzeba co najmniej dwóch dodatnich czasów.')
        T=float(times.mean()/n);uA=float(times.std(ddof=1)/math.sqrt(len(times))/n)
        ut=uncertainty(p,'u_t_s');uL=uncertainty(p,'u_L_m')
        g=4*math.pi**2*L/T**2;uT=math.hypot(uA,ut/n) if ut is not None else None
        ug=g*math.hypot(uL/L,2*uT/T) if uL is not None and uT is not None else None
        r.results['stala_dlugosc']=dict(n=len(times),T_s=T,uA_T_s=uA,g_m_s2=g,u_g_m_s2=ug)
        r.heading('Stała długość');r.table(['T [s]','uA(T) [s]','g [m/s²]','u(g) [m/s²]'],[[T,uA,g,ug if ug is not None else 'brak pełnych danych']])
        if ug is None:r.warn('Własna seria: pełna u(g) wymaga u(L) i u pojedynczego czasu. Wyświetlona uA(T) obejmuje tylko rozrzut serii.')
        axes('Numer pomiaru','Okres T [s]');plt.plot(range(1,len(times)+1),times/n,'o-',color='#4472c4');r.savefig('okresy','Okres jednego drgania w kolejnych powtórzeniach.')
    if t['Grupy']:
        rows=t['Grupy'];x=values(rows,'L_cm')*.01;T=values(rows,'T_s');ux=values(rows,'u_L_cm')*.01;ut=values(rows,'u_T_s')
        if np.any(x<=0) or np.any(T<=0) or np.any(ux<0) or np.any(ut<0):raise ValueError('Grupy: dodatnie L,T oraz nieujemne niepewności.')
        a,b,ua,res=fit(x,T*T,True);g=4*math.pi**2/a
        if a<=0:raise ValueError('Niedodatnie nachylenie T²(L).')
        r.heading('Metoda T²(L)');r.table(['A [s²/m]','u dopasowania A','g [m/s²]','u dopasowania g'],[[a,ua,g,g*ua/a]])
        r.results['regresja']=dict(A=a,u_A=ua,g=g,u_g_fit=g*ua/a)
        fig,ax=axes('Długość L [m]','Kwadrat okresu T² [s²]');ax.errorbar(x,T*T,xerr=ux,yerr=2*T*ut,fmt='o',capsize=3,label='Pomiar');xx=np.linspace(0,x.max(),200);ax.plot(xx,a*xx,label='OLS przez początek');ax.legend();r.savefig('T2','Kwadrat okresu od długości; słupki oznaczają wpisane niepewności standardowe.')
        r.warn('Regresja T²(L) jest OLS przez początek układu. u dopasowania wynika z reszt i nie zastępuje pełnej analizy wpływu niepewności długości i okresu.')


def young(r,t):
    p=r.p;g=require(p,'g',True);samples={str(x['nr']):x for x in t['Probki']}
    for key in ['A','B']:
        rows=t['Drut_'+key]
        if not rows:continue
        if key not in samples:raise ValueError('Uzupełnij wymiary drutu '+key)
        s=samples[key];L=num(s['l_m'],'l drutu '+key,positive=True)
        d=np.mean([num(s[k],k+' drutu '+key,positive=True) for k in ['d1_mm','d2_mm','d3_mm']])*.001
        z0u=num(s['z0_up_mm'],'zero obciążania '+key);z0d=num(s['z0_down_mm'],'zero odciążania '+key)
        mass=values(rows,'m_kg');up=values(rows,'z_up_mm');down=values(rows,'z_down_mm')
        if np.any(mass<0):raise ValueError('Masa nie może być ujemna.')
        material=str(s.get('material') or key)
        if ('stal' in material.lower() and mass.max()>10) or (any(z in material.lower() for z in ['mosi','mied']) and mass.max()>6):
            raise ValueError('Masa przekracza limit instrukcji dla drutu '+key)
        select=[]
        for row in rows:
            flag=str(row.get('fit') or 'tak').lower().strip()
            if flag not in ['tak','nie']:raise ValueError('Regresja: wpisz tak albo nie.')
            select.append(flag=='tak')
        select=np.array(select);force=mass*g;elong=(up-z0u+down-z0d)/4*.001
        a,b,ua,res=fit(force[select],elong[select])
        if a<=0:raise ValueError('Drut '+key+': nachylenie wydłużenia musi być dodatnie. Sprawdź zera i odczyty.')
        E=4*L/(math.pi*d*d*a);uL=uncertainty(s,'u_l_m');ud=uncertainty(s,'u_d_mm')
        uE=E*math.sqrt((uL/L)**2+(2*ud*.001/d)**2+(ua/a)**2) if uL is not None and ud is not None else None
        r.results[key]=dict(a_m_N=a,b_m=b,u_a=ua,E_GPa=E/1e9,u_E_GPa=uE/1e9 if uE is not None else None,n_fit=int(select.sum()))
        r.heading('Drut '+key+' — '+material)
        r.table(['l [m]','d1 [mm]','d2 [mm]','d3 [mm]','Średnia d [mm]'],[[L,s['d1_mm'],s['d2_mm'],s['d3_mm'],d*1000]])
        r.text(f'Zero obciążania: {z0u:g} mm; zero odciążania: {z0d:g} mm. u(l): '+(f'{uL:g} m' if uL is not None else 'nie podano')+'; u(d): '+(f'{ud:g} mm.' if ud is not None else 'nie podano.'))
        r.table(['a [m/N]','b [m]','E [GPa]','u(E) [GPa]'],[[a,b,E/1e9,uE/1e9 if uE is not None else 'brak u(l),u(d)']])
        r.table(['m [kg]','z↑ [mm]','z↓ [mm]','Δl [mm]','Regresja'],[[m,u,v,e*1000,'tak' if flag else 'nie'] for m,u,v,e,flag in zip(mass,up,down,elong,select)])
        fig,ax=axes('Siła F [N]','Wydłużenie Δl [mm]');ax.plot(force,elong*1000,'o',label='Wszystkie pomiary');xx=np.linspace(force.min(),force.max(),200);ax.plot(xx,(a*xx+b)*1000,label='Regresja wybranego zakresu');ax.legend();r.savefig('drut_'+key,'Wydłużenie drutu i regresja ze swobodnym wyrazem wolnym.')
        if uE is None:r.warn('Drut '+key+': bez u(l) i u(d) nie wyznaczono pełnej u(E).')
    r.warn('Niepewność regresji Δl(F) obejmuje rozrzut reszt. Model przyjmuje pomijalną niepewność siły; podaj uzasadnienie i wpływ dokładności odważników.')


def field(r,t):
    p=r.p;variant=r.spec['code'][-1]
    if t.get('Profile'):
        rows=t['Profile'];x=values(rows,'s_mm');V=np.column_stack([values(rows,k) for k in ['Va','Vb','Vc']])
        order=np.argsort(x);x=x[order];V=V[order]
        if len(x)<2 or np.any(np.diff(x)<=0):raise ValueError('Profile: co najmniej 2 różne położenia, bez duplikatów.')
        U=require(p,'U');V0=require(p,'V0');mean=V.mean(axis=1);mid=(x[:-1]+x[1:])/2;E=-np.diff(mean)/(np.diff(x)*.001)
        if variant=='A':
            d=require(p,'d_mm',True);theory=V0+U*x/d;Et=np.full(len(mid),-U/(d*.001))
            if np.any((x<0)|(x>d)):raise ValueError('Profile A muszą leżeć między elektrodami: 0 ≤ x ≤ d.')
        else:
            rw=require(p,'rw_mm',True);rz=require(p,'rz_mm',True)
            if rz<=rw or np.any((x<rw)|(x>rz)):raise ValueError('B: wymagane rw < rz i rw ≤ r ≤ rz.')
            theory=V0+U*np.log(rz/x)/np.log(rz/rw);Et=U/(mid*.001*np.log(rz/rw))
        r.results['profile']=dict(s_mm=x.tolist(),mean_V=mean.tolist(),mid_mm=mid.tolist(),E_V_m=E.tolist(),theory_V=theory.tolist(),theory_E=Et.tolist())
        uV=uncertainty(p,'u_V');us=uncertainty(p,'u_s_mm')
        uE=None
        if uV is not None and us is not None:
            ds=np.diff(x)*.001
            uE=np.sqrt(2*uV**2/ds**2+np.diff(mean)**2*(us*.001)**2/ds**4)
            r.results['profile']['u_E_V_m']=uE.tolist()
            r.warn('u(E) w profilu przyjmuje wpisaną u(V) dla każdego średniego potencjału, bez redukcji przez √3, oraz u odstępu Δs. Wymaga niezależności sąsiednich potencjałów i odstępu. Wspólny błąd odniesienia może zmienić tę propagację.')
        r.heading('Profile potencjału i natężenia');r.table(['s [mm]','Va [V]','Vb [V]','Vc [V]','Średnia [V]','Teoria [V]'],[[xx,*vv,vm,vt] for xx,vv,vm,vt in zip(x,V,mean,theory)])
        r.table(['s* [mm]','E ze znakiem [V/m]','E teor. [V/m]','u(E) [V/m]'],
                [[ss,ee,et,float(uE[i]) if uE is not None else 'brak u(V),u(Δs)'] for i,(ss,ee,et) in enumerate(zip(mid,E,Et))])
        fig,ax=axes('x [mm]' if variant=='A' else 'r [mm]','Potencjał V [V]');ax.plot(x,mean,'o-',label='Średnia pomiarów');ax.plot(x,theory,'o-',label='Teoria');ax.legend();r.savefig('potencjal','Średnia trzech profili oraz potencjał teoretyczny dla wpisanych napięć.')
        fig,ax=axes('Położenie środka przedziału [mm]','Natężenie |E| [V/m]');ax.plot(mid,abs(E),'o-',label='Pomiar');ax.plot(mid,abs(Et),'o-',label='Teoria');ax.legend();r.savefig('pole','Moduł pola; tabela zachowuje znak składowej. Odcinki łączą kolejne punkty.')
        r.warn('Trzy kierunki profilu nie są automatycznie niezależnymi powtórzeniami. Niepewności różnic potencjału zależą również od wspólnego odniesienia i kroku położenia.')
    if t.get('Mapa'):
        rows=t['Mapa'];x=values(rows,'x_mm');y=values(rows,'y_mm');v=values(rows,'V')
        if len(set(zip(x,y)))!=len(x):raise ValueError('Mapa: powtórzone współrzędne. Powtórzenia opracuj osobno.')
        r.heading('Mapa potencjału');fig,ax=axes('x [mm]','y [mm]');ax.set_aspect('equal');pts=ax.scatter(x,y,c=v,cmap='viridis',zorder=3);fig.colorbar(pts,ax=ax,label='Potencjał [V]')
        h=uncertainty(p,'h_mm');k=uncertainty(p,'k_mm');vectors=[]
        if h is not None and k is not None and h>0 and k>0 and len(x)>=3 and np.linalg.matrix_rank(np.column_stack([x-x[0],y-y[0]]))==2:
            mesh=tri.Triangulation(x,y);coords=np.column_stack([x,y])[mesh.triangles]
            edge=np.linalg.norm(coords-np.roll(coords,1,axis=1),axis=2);mask=edge.max(axis=1)>math.hypot(h,k)*1.01
            mesh.set_mask(mask)
            if np.any(~mask):
                ax.tricontour(mesh,v,levels=8,colors='#e83d58',linewidths=.9)
            points={(round(xx,8),round(yy,8)):vv for xx,yy,vv in zip(x,y,v)}
            for (xx,yy),vv in points.items():
                right=points.get((round(xx+h,8),yy));top=points.get((xx,round(yy+k,8)))
                if right is not None and top is not None:vectors.append([xx,yy,-(right-vv)/(h*.001),-(top-vv)/(k*.001)])
            if vectors:
                a=np.array(vectors);ax.quiver(a[:,0],a[:,1],a[:,2],a[:,3],color='#1765b3',angles='xy')
        r.savefig('mapa','Punkty potencjału; kontury tylko w lokalnych trójkątach o rozmiarze zgodnym z h,k. Strzałki pokazują kierunek ujemnego gradientu; długości są skalowane dla czytelności.')
        r.results['mapa']=dict(n=len(rows),wektory=vectors)
        if vectors:r.table(['x [mm]','y [mm]','Ex [V/m]','Ey [V/m]'],vectors)
        else:r.warn('Mapa: nie wyznaczono wektorów. Potrzeba dodatnich h,k i zmierzonych sąsiadów (x+h,y), (x,y+k).')
        r.warn('Mapa nie odtwarza potencjału poza pomiarami. Kontury są interpolacją, a składowe wektorów różnicami w przód w punkcie bazowym. Długość strzałki nie jest wzorcem V/m.')


def slits(r,t):
    p=r.p
    if t['Pomiary']:
        rows=t['Pomiary'];x=values(rows,'p_mm');u=values(rows,'U_mV')
        if len(set(x))!=len(x):raise ValueError('Powtórzone położenia w tabeli Pomiary.')
        order=np.argsort(x);x=x[order];u=u[order]
        r.heading('Rozkład sygnału');r.text(f'Liczba odczytów: {len(x)}. Zakres położenia: {x.min():g}–{x.max():g} mm. Nie odejmowano automatycznie tła.')
        axes('Położenie p [mm]','Sygnał U [mV]');plt.plot(x,u,'o-',ms=2,lw=.8);r.savefig('sygnal','Surowe napięcia detektora. Położenie pochodzi z odczytu lub jawnego początku i kroku w Excelu.')
        r.results['pomiary']=dict(n=len(x),p_min=float(x.min()),p_max=float(x.max()))
    if t['Maksima']:
        wave=require(p,'lambda_nm',True)*1e-6;L=require(p,'L_mm',True);results=[];orders=[]
        u_wave=uncertainty(p,'u_lambda_nm');uL=uncertainty(p,'u_L_mm');up=uncertainty(p,'u_p_mm')
        for row in t['Maksima']:
            m=num(row['m'],'rząd maksimum',positive=True);left=num(row['left_mm'],'lewe maksimum');right=num(row['right_mm'],'prawe maksimum')
            if m!=int(m) or right<=left or m in orders:raise ValueError('Maksima: dodatnie, różne całkowite rzędy i lewe położenie mniejsze od prawego.')
            orders.append(m);half=(right-left)/2;d=m*wave*L/half
            ud=d*math.sqrt((u_wave*1e-6/wave)**2+(uL/L)**2+(up/math.sqrt(2)/half)**2) if None not in [u_wave,uL,up] else None
            results.append([m,left,right,half,d,ud])
        ds=np.array([x[4] for x in results]);ua=float(ds.std(ddof=1)/math.sqrt(len(ds))) if len(ds)>1 else None
        r.heading('Rozstaw szczelin');r.table(['m','Lewe [mm]','Prawe [mm]','r [mm]','d [mm]','u(d) [mm]'],[[*a[:5],a[5] if a[5] is not None else 'brak danych'] for a in results])
        r.text(f'Średni rozstaw d = {ds.mean():.6g} mm.'+(f' Składowa typu A: {ua:.4g} mm.' if ua is not None else ' Jeden wynik nie pozwala wyznaczyć uA średniej.'))
        r.results['szczeliny']=dict(d_mm=float(ds.mean()),uA_mm=ua,pary=results)
        axes('Rząd m','Połówka rozstawu r [mm]');plt.plot(orders,[x[3] for x in results],'o');r.savefig('maksima','Połówki rozstawu par maksimów według wpisanych rzędów.')
        r.warn('uA średniej rozstawu nie jest pełną niepewnością końcową. λ i L są wspólne dla par, więc ich wkład nie maleje jak 1/√N. Wzór d=mλL/r korzysta z przybliżenia małych kątów.')
    elif t['Pomiary']:r.warn('Aby obliczyć rozstaw d, uzupełnij λ, L oraz tabelę Maksima. Program nie zgaduje numerów prążków.')
    if not missing(p.get('Imax')) and not missing(p.get('Imin')):
        hi=require(p,'Imax',True);lo=require(p,'Imin')
        if not 0<=lo<=hi:raise ValueError('Wymagane 0 ≤ Imin ≤ Imax.')
        r.table(['Q = Imin/Imax','Widzialność'],[[lo/hi,(hi-lo)/(hi+lo)]])
        r.results['kontrast']=dict(Q=lo/hi,widzialnosc=(hi-lo)/(hi+lo))


def spectrum(r,t):
    p=r.p;N=require(p,'N',True);d=1e6/N;groups={};seen=set()
    refs={}
    for row in t['Odniesienia']:
        idn=num(row['id'],'ID odniesienia',positive=True)
        if idn!=int(idn) or int(idn) in refs:raise ValueError('Odniesienia: ID musi być całkowite i nie może się powtarzać.')
        refs[int(idn)]=row
    for row in t['Odczyty']:
        order=num(row['order'],'rząd',positive=True);idn=num(row['id'],'ID linii',positive=True);repeat=num(row['repeat'],'powtórzenie',positive=True)
        if order not in [1,2] or idn!=int(idn) or repeat!=int(repeat):raise ValueError('Widmo: całkowite ID i powtórzenie oraz rząd 1 albo 2.')
        key=(int(idn),int(order));tag=(*key,int(repeat))
        if tag in seen:raise ValueError('Powtórzone ID/rząd/powtórzenie w Odczytach.')
        seen.add(tag)
        def angle(deg,minute):
            D=num(deg,'stopnie',nonnegative=True);M=num(minute,'minuty',nonnegative=True)
            if D>=360 or D!=int(D) or M>=60:raise ValueError('Stopnie: liczba całkowita 0–359, minuty: 0 ≤ M < 60.')
            return D+M/60
        zero=angle(p.get(f'zero{int(order)}_deg'),p.get(f'zero{int(order)}_min'))
        left=angle(row['ldeg'],row['lmin']);right=angle(row['rdeg'],row['rmin'])
        signed_l=(left-zero+180)%360-180;signed_r=(right-zero+180)%360-180
        al=abs(signed_l);ar=abs(signed_r)
        if max(al,ar)>=90:raise ValueError('Odchylenie ≥90°: sprawdź kierunek zerowy i odczyty.')
        zero_coefficient=-(np.sign(signed_l)+np.sign(signed_r))/2
        groups.setdefault(key,[]).append((al,ar,str(row.get('barwa') or ''),float(zero_coefficient)))
    results=[]
    for (idn,order),vv in sorted(groups.items()):
        al=float(np.mean([v[0] for v in vv]));ar=float(np.mean([v[1] for v in vv]));alpha=(al+ar)/2
        wave=d*math.sin(math.radians(alpha))/order
        u_angle=uncertainty(p,'u_angle_min');uN=uncertainty(p,'u_N')
        # Zachowujemy korelację wspólnego zera. Dla stron po przeciwnych
        # stronach θ0 jego wkład w (αL+αP)/2 znosi się algebraicznie.
        zero_coefficient=float(np.mean([v[3] for v in vv]))
        ua=math.radians(u_angle/60)*math.sqrt(zero_coefficient**2+1/(2*len(vv))) if u_angle is not None else None
        uw=math.hypot(d*math.cos(math.radians(alpha))/order*ua,wave*uN/N) if ua is not None and uN is not None else None
        results.append(dict(id=idn,order=order,n=len(vv),barwa=vv[0][2],al=al,ar=ar,alpha=alpha,lambda_nm=wave,u_lambda_nm=uw))
    if not results:raise ValueError('Brak odczytów kątowych w Odczytach.')
    r.heading('Opracowanie obu stron i rzędów');r.table(['ID','Rząd','αL [°]','αP [°]','λ [nm]','u(λ) [nm]'],[[v['id'],v['order'],v['al'],v['ar'],v['lambda_nm'],v['u_lambda_nm'] if v['u_lambda_nm'] is not None else 'brak danych'] for v in results])
    final=[]
    for idn in sorted({v['id'] for v in results}):
        lines=[v for v in results if v['id']==idn];mean=float(np.mean([v['lambda_nm'] for v in lines]));ref=refs.get(idn,{})
        reference=None if missing(ref.get('lambda_ref')) else num(ref['lambda_ref'],'λ odniesienia',positive=True)
        final.append([idn,lines[0]['barwa'],mean,reference if reference is not None else '—',mean-reference if reference is not None else '—',ref.get('przejscie') or '—'])
    r.table(['ID','Barwa','Średnia λ [nm]','Odniesienie [nm]','Różnica [nm]','Przejście'],final)
    r.results['widmo']=results;r.results['linie_koncowe']=final
    fig,ax=axes('ID linii','Długość fali λ [nm]')
    for order in [1,2]:
        vv=[v for v in results if v['order']==order]
        if vv:ax.plot([v['id'] for v in vv],[v['lambda_nm'] for v in vv],'o',label=f'Rząd {order}')
    ax.legend();r.savefig('widmo','Długości fal wyznaczone oddzielnie w I i II rzędzie.')
    r.warn('Średnia długości fal z rzędów nie jest niezależnym powtórzeniem: wspólne są stała siatki i kierunki zerowe. Podane u(λ) dotyczą pojedynczego rzędu, bez dodatkowej oceny rozrzutu powtórzeń. Nie podano automatycznie niepewności końcowej średniej.')
    r.warn('Model u(λ) zakłada niezależne wskazania stron, wspólny odczyt θ0 w obrębie rzędu i tę samą standardową u kąta dla każdego wskazania. Korelację zera zachowano: dla przeciwnych stron jego wkład w średnie odchylenie znosi się. Nie obejmuje to np. systematycznego błędu skali kątowej.')
    if any(v['u_lambda_nm'] is None for v in results):r.warn('Brak u kąta lub u(N): nie wyznaczono u(λ). Brak danych nie oznacza zerowej niepewności.')
    if not missing(p.get('kontrola_l')) and not missing(p.get('kontrola_p')):
        diff=abs(require(p,'kontrola_l')-require(p,'kontrola_p'))*60;r.text(f'Kontrola prostopadłości: różnica odchyleń {diff:.3g} min.')
        if diff>10:r.warn('Różnica stron linii kontrolnej przekracza 10 minut kątowych; sprawdź prostopadłość siatki.')


def generate(source,output=None,tectonic=None,tex_only=False):
    spec,p,t=load(source)
    if not any(t.values()):raise ValueError('Brak pomiarów: wypełnij co najmniej jedną tabelę w Excelu.')
    out=Path(output) if output else Path(source).resolve().parent/'Wyniki'/spec['code']
    out.mkdir(parents=True,exist_ok=True)
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':11,'axes.spines.top':False,'axes.spines.right':False,'pdf.fonttype':42})
    r=Report(spec,p,out)
    if spec['code']=='0':pendulum(r,t)
    elif spec['code']=='11':young(r,t)
    elif spec['code'].startswith('31'):field(r,t)
    elif spec['code']=='71':slits(r,t)
    else:spectrum(r,t)
    if not r.results:raise ValueError('Nie ma kompletnej serii pomiarowej do opracowania.')
    r.write()
    for name,rows in t.items():
        if not rows:continue
        with (out/(name+'.csv')).open('w',encoding='utf-8-sig',newline='') as f:
            keys=[k for k in rows[0] if k!='_row'];writer=csv.DictWriter(f,keys,extrasaction='ignore');writer.writeheader();writer.writerows(rows)
    # Paczka zawiera tylko rysunki bieżącego raportu, bez źródłowych danych XLSX.
    with zipfile.ZipFile(out/'sprawozdanie_overleaf.zip','w',zipfile.ZIP_DEFLATED) as z:
        z.write(out/'sprawozdanie.tex','sprawozdanie.tex')
        import re
        for name in re.findall(r'\\includegraphics\[[^]]*\]\{([^}]+)\}',(out/'sprawozdanie.tex').read_text(encoding='utf-8')):
            z.write(out/name,name)
    engine=tectonic or os.environ.get('FIZLAB_TECTONIC') or shutil.which('tectonic')
    pdf=None
    if not tex_only and engine:
        target=out/'sprawozdanie.pdf'
        # Stary PDF nie może udawać nowego wyniku przy nieudanej kompilacji.
        if target.exists():target.rename(out/'sprawozdanie_poprzednie.pdf') if not (out/'sprawozdanie_poprzednie.pdf').exists() else target.unlink()
        process=subprocess.run([str(engine),'sprawozdanie.tex'],cwd=out,capture_output=True,text=True,encoding='utf-8',errors='replace')
        (out/'kompilacja.log').write_text(process.stdout+'\n'+process.stderr,encoding='utf-8')
        if process.returncode:raise ValueError('Kompilacja PDF nie powiodła się; sprawdź kompilacja.log. Źródła i ZIP zostały zachowane.')
        pdf=target
    return r,out,pdf


def main():
    if hasattr(sys.stdout,'reconfigure'):sys.stdout.reconfigure(encoding='utf-8')
    if hasattr(sys.stderr,'reconfigure'):sys.stderr.reconfigure(encoding='utf-8')
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('excel',type=Path);parser.add_argument('--output-dir',type=Path)
    parser.add_argument('--tectonic',help='Ścieżka do programu Tectonic; domyślnie PATH lub FIZLAB_TECTONIC.')
    parser.add_argument('--tex-only',action='store_true',help='Tylko LaTeX i ZIP do Overleaf.')
    args=parser.parse_args()
    try:r,out,pdf=generate(args.excel,args.output_dir,args.tectonic,args.tex_only)
    except (ValueError,KeyError,OSError,zipfile.BadZipFile) as e:parser.exit(2,f'Błąd: {e}\n')
    print('Zapisano:',out)
    print('PDF:',pdf if pdf else 'Nie kompilowano. Otwórz sprawozdanie_overleaf.zip w Overleaf (XeLaTeX) lub zainstaluj Tectonic.')
    for w in dict.fromkeys(r.warnings):print('Uwaga:',w)


if __name__=='__main__':main()
